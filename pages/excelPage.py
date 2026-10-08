import datetime
import os

from openpyxl import load_workbook
from PySide6 import QtCore, QtWidgets


class excelPage(QtWidgets.QWidget):
    def __init__(self, pageSwitchFunction):
        super().__init__()
        self.pageSwitch = pageSwitchFunction
        self._sheets = {}

        self.back_button = QtWidgets.QPushButton("← Takaisin")
        self.back_button.setObjectName("backButton")
        self.back_button.setCursor(QtCore.Qt.PointingHandCursor)
        self.back_button.clicked.connect(lambda: self.pageSwitch(0))

        self.file_label = QtWidgets.QLabel("Ei tiedostoa valittu")
        self.file_label.setProperty("role", "fileLabel")

        self.sheet_combo = QtWidgets.QComboBox()
        self.sheet_combo.setVisible(False)
        self.sheet_combo.currentTextChanged.connect(self._show_sheet)

        self.table = QtWidgets.QTableWidget()
        self.table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)

        top = QtWidgets.QHBoxLayout()
        top.setContentsMargins(12, 12, 12, 4)
        top.addWidget(self.back_button)
        top.addSpacing(12)
        top.addWidget(self.file_label, 1)
        top.addWidget(self.sheet_combo)

        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(16, 8, 16, 16)
        layout.addLayout(top)
        layout.addWidget(self.table, 1)

    def open_excel_file(self):
        filepath, _ = QtWidgets.QFileDialog.getOpenFileName(
            self,
            "Avaa Excel-tiedosto",
            "",
            "Excel Files (*.xlsx *.xlsm)",
        )
        if not filepath:
            return False
        return self.load_excel(filepath)

    def load_excel(self, filepath):
        try:
            workbook = load_workbook(filepath, data_only=True, read_only=True)
        except Exception:
            self.file_label.setText("Excel-tiedostoa ei voitu avata")
            return False

        sheets = {}
        try:
            for name in workbook.sheetnames:
                sheets[name] = [
                    [self._cell_text(cell) for cell in row]
                    for row in workbook[name].iter_rows(values_only=True)
                ]
        except Exception:
            self.file_label.setText("Excel-tiedostoa ei voitu lukea")
            return False
        finally:
            workbook.close()

        if not sheets:
            self.file_label.setText("Excel-tiedostossa ei ole taulukoita")
            return False

        self._sheets = sheets
        self._filepath = filepath

        self.sheet_combo.blockSignals(True)
        self.sheet_combo.clear()
        self.sheet_combo.addItems(list(sheets.keys()))
        self.sheet_combo.setVisible(len(sheets) > 1)
        self.sheet_combo.blockSignals(False)

        self._show_sheet(self.sheet_combo.currentText())
        return True

    def _show_sheet(self, sheet_name):
        rows = self._sheets.get(sheet_name) or []
        filename = os.path.basename(getattr(self, "_filepath", ""))
        self._fill_table(rows, filename, sheet_name)

    def _fill_table(self, rows, filename, sheet_name):
        rows = [row for row in rows if any(cell != "" for cell in row)]
        if not rows:
            self.table.clear()
            self.table.setRowCount(0)
            self.table.setColumnCount(0)
            self.file_label.setText(f"{filename} — tyhjä taulukko")
            return

        headers = rows[0]
        data = rows[1:]
        column_count = max((len(row) for row in rows), default=0)
        headers = list(headers) + [""] * (column_count - len(headers))

        self.table.clear()
        self.table.setColumnCount(column_count)
        self.table.setRowCount(len(data))
        self.table.setHorizontalHeaderLabels(headers)

        for row_index, row in enumerate(data):
            for column_index in range(column_count):
                value = row[column_index] if column_index < len(row) else ""
                self.table.setItem(
                    row_index,
                    column_index,
                    QtWidgets.QTableWidgetItem(value),
                )

        self.table.resizeColumnsToContents()
        sheet_note = f" / {sheet_name}" if len(self._sheets) > 1 else ""
        self.file_label.setText(f"{filename}{sheet_note} — {len(data)} riviä")

    def _cell_text(self, value):
        if value is None:
            return ""
        if isinstance(value, datetime.datetime):
            if value.time() == datetime.time():
                return value.strftime("%d.%m.%Y")
            return value.strftime("%d.%m.%Y %H:%M")
        if isinstance(value, datetime.date):
            return value.strftime("%d.%m.%Y")
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
